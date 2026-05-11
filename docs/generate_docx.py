#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
毕业论文Markdown转Word工具
根据湖南农业大学毕业论文格式要求生成Word文档
"""

import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn


class ThesisDocxGenerator:
    """论文Word文档生成器"""
    
    def __init__(self):
        self.doc = Document()
        self.setup_page_layout()
        
    def setup_page_layout(self):
        """设置页面布局"""
        for section in self.doc.sections:
            section.top_margin = Cm(2.54)
            section.bottom_margin = Cm(2.54)
            section.left_margin = Cm(2.6)
            section.right_margin = Cm(2.6)
    
    def add_title(self, text, level=1):
        """添加标题"""
        p = self.doc.add_paragraph()
        
        if level == 1:
            # 一级标题：小三号黑体（15磅）
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            run.font.size = Pt(15)
            run.bold = True
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
        elif level == 2:
            # 二级标题：四号黑体（14磅）
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            run.font.size = Pt(14)
            run.bold = True
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
        elif level == 3:
            # 三级标题：小四号黑体（12磅）
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
            run.font.size = Pt(12)
            run.bold = True
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
        
        # 固定行距22磅
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(22)
        
    def add_paragraph(self, text, indent=True, alignment=None):
        """添加段落"""
        if not text.strip():
            return
            
        p = self.doc.add_paragraph()
        
        # 固定行距22磅
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(22)
        
        # 首行缩进
        if indent:
            p.paragraph_format.first_line_indent = Pt(24)  # 2字符
        
        if alignment:
            p.alignment = alignment
        
        # 普通段落（小四号宋体）
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(12)
    
    def add_reference(self, text):
        """添加参考文献条目"""
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(22)
        
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(10.5)  # 五号字
        
    def add_cover_title(self, text):
        """添加正文开始的论文标题（小二黑体，居中）"""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        run.font.size = Pt(18)  # 小二号
        run.bold = True
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(22)
        
    def add_author_info(self, text):
        """添加学生/指导老师信息（五号宋体，居中）"""
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
        run.font.size = Pt(10.5)  # 五号
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.line_spacing = Pt(22)
    
    def parse_and_generate(self, md_file, output_file):
        """解析Markdown文件并生成Word文档"""
        with open(md_file, 'r', encoding='utf-8') as f:
            lines = f.readlines()
        
        i = 0
        in_cover_section = True  # 标记是否在封面和诚信声明部分
        found_main_title = False  # 标记是否找到正文标题
        
        while i < len(lines):
            line = lines[i].rstrip()
            
            # 跳过空行
            if not line:
                i += 1
                continue
            
            # 检测是否到达正文开始（通过检测"基于大语言模型"这样的论文标题）
            if in_cover_section and '基于大语言模型' in line and '学　　生' not in line:
                in_cover_section = False
                if not found_main_title and not line.startswith('摘'):
                    # 这是正文的论文标题
                    self.add_cover_title(line)
                    found_main_title = True
                    i += 1
                    continue
            
            # 在封面和诚信声明部分，跳过处理（让用户在Word中手动调整）
            if in_cover_section:
                i += 1
                continue
            
            # 处理三级标题 (x.x.x 格式)
            if re.match(r'^\d+\.\d+\.\d+\s+', line):
                self.add_title(line, level=3)
                i += 1
                continue
            
            # 处理二级标题 (x.x 格式)
            if re.match(r'^\d+\.\d+\s+', line):
                self.add_title(line, level=2)
                i += 1
                continue
            
            # 处理一级标题 (x 格式)
            if re.match(r'^\d+\s+[^\d]', line):
                self.add_title(line, level=1)
                i += 1
                continue
            
            # 处理学生信息行（五号宋体，居中）
            if line.startswith('学　　生：') or line.startswith('指导老师：') or '(湖南农业大学' in line:
                self.add_author_info(line)
                i += 1
                continue
            
            # 处理Student/Tutor行（五号Times New Roman，居中）
            if line.startswith('Student:') or line.startswith('Tutor:') or '(College of' in line:
                self.add_author_info(line)
                i += 1
                continue
            
            # 处理摘要（标题加粗，内容五号字）
            if line.startswith('摘　要：') or line.startswith('摘要：'):
                content = line.replace('摘　要：', '').replace('摘要：', '')
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(22)
                p.paragraph_format.first_line_indent = Pt(24)  # 2字符
                # 粗体部分（小四黑体）
                run1 = p.add_run('摘　要：')
                run1.font.name = 'Times New Roman'
                run1._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run1.font.size = Pt(12)  # 小四号
                run1.bold = True
                # 正文部分（五号宋体）
                run2 = p.add_run(content)
                run2.font.name = 'Times New Roman'
                run2._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                run2.font.size = Pt(10.5)  # 五号
                i += 1
                continue
            
            # 处理关键词（标题加粗，内容五号字）
            if line.startswith('关键词：'):
                content = line.replace('关键词：', '')
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(22)
                p.paragraph_format.first_line_indent = Pt(24)  # 2字符
                # 粗体部分（小四黑体）
                run1 = p.add_run('关键词：')
                run1.font.name = 'Times New Roman'
                run1._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run1.font.size = Pt(12)  # 小四号
                run1.bold = True
                # 正文部分（五号宋体）
                run2 = p.add_run(content)
                run2.font.name = 'Times New Roman'
                run2._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
                run2.font.size = Pt(10.5)  # 五号
                i += 1
                continue
            
            # 处理Abstract（标题加粗，内容五号Times New Roman）
            if line.startswith('Abstract:') or line.startswith('Abstract：'):
                content = line.replace('Abstract:', '').replace('Abstract：', '').strip()
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(22)
                p.paragraph_format.first_line_indent = Pt(24)  # 2字符
                # 粗体部分（小四号Times New Roman加粗）
                run1 = p.add_run('Abstract: ')
                run1.font.name = 'Times New Roman'
                run1.font.size = Pt(12)  # 小四号
                run1.bold = True
                # 正文部分（五号Times New Roman）
                run2 = p.add_run(content)
                run2.font.name = 'Times New Roman'
                run2.font.size = Pt(10.5)  # 五号
                i += 1
                continue
            
            # 处理Key words（标题加粗，内容五号Times New Roman）
            if line.startswith('Key words:'):
                content = line.replace('Key words:', '').strip()
                p = self.doc.add_paragraph()
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(22)
                p.paragraph_format.first_line_indent = Pt(24)  # 2字符
                # 粗体部分（小四号Times New Roman加粗）
                run1 = p.add_run('Key words: ')
                run1.font.name = 'Times New Roman'
                run1.font.size = Pt(12)  # 小四号
                run1.bold = True
                # 正文部分（五号Times New Roman）
                run2 = p.add_run(content)
                run2.font.name = 'Times New Roman'
                run2.font.size = Pt(10.5)  # 五号
                i += 1
                continue
            
            # 处理参考文献标题
            if line == '参考文献':
                self.add_title('参考文献', level=1)
                i += 1
                continue
            
            # 处理参考文献条目
            if re.match(r'^\[\d+\]', line):
                self.add_reference(line)
                i += 1
                continue
            
            # 处理致谢标题（居中）
            if line == '致　　谢':
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run('致　　谢')
                run.font.name = 'Times New Roman'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
                run.font.size = Pt(15)  # 小三号
                run.bold = True
                p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                p.paragraph_format.line_spacing = Pt(22)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                i += 1
                continue
            
            # 跳过图表占位符
            if '【此处插入图' in line or '【此处插入表' in line:
                i += 1
                continue
            
            # 跳过图表说明行（Fig.x 或 Table.x）
            if re.match(r'^Fig\.\d+|^Table\s*\d+', line):
                i += 1
                continue
            
            # 跳过以"图x"或"表x"开头的行
            if re.match(r'^图\d+|^表\d+', line):
                i += 1
                continue
            
            # 普通段落
            # 跳过过短的行（可能是标题或格式行）
            if len(line.strip()) > 5:
                self.add_paragraph(line)
            
            i += 1
        
        # 保存文档
        self.doc.save(output_file)
        print(f"✅ Word文档已生成: {output_file}")


def main():
    """主函数"""
    import os
    from datetime import datetime
    
    # 获取当前脚本所在目录
    current_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(current_dir, '毕业论文_初稿.md')
    # 使用时间戳生成唯一文件名
    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    output_file = os.path.join(current_dir, f'毕业论文_格式化版_{timestamp}.docx')
    
    generator = ThesisDocxGenerator()
    generator.parse_and_generate(md_file, output_file)


if __name__ == '__main__':
    main()
